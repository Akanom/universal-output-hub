from pathlib import Path

import pandas as pd

from universal_output_hub import OutputHub


def test_dictionary_model_and_exports(tmp_path: Path) -> None:
    hub = OutputHub("Test")
    hub.add_model(
        {
            "name": "M1",
            "params": {"x": 1.0, "z": -2.0},
            "std_errors": {"x": 0.1, "z": 0.2},
            "pvalues": {"x": 0.01, "z": 0.20},
            "statistics": {"N": 100},
            "diagnostics": {"Hansen p": 0.50},
        }
    )
    table = hub.regression_table()
    assert "M1" in table.columns
    assert "x" in table.index
    assert table.loc["x", "M1"] == "1.000***"
    assert table.loc["z", "M1"] == "-2.000"
    assert "N" in table.index
    assert "Significance" in table.index
    bundle = hub.export_bundle(tmp_path)
    assert bundle["manifest"].exists()
    assert bundle["html_report"].exists()


def test_external_coefficient_table() -> None:
    hub = OutputHub("External")
    df = pd.DataFrame({"term": ["x"], "coef": [1.2], "se": [0.3], "pvalue": [0.04]})
    hub.add_model_table(df, name="External Model")
    out = hub.regression_table()
    assert out.loc["x", "External Model"] == "1.200**"


def test_custom_star_levels_and_no_stars() -> None:
    hub = OutputHub("Stars")
    hub.add_model({"name": "M", "params": {"x": 1.0}, "std_errors": {"x": 0.1}, "pvalues": {"x": 0.04}})

    custom = hub.regression_table(star_levels={"***": 0.001, "**": 0.01, "*": 0.05})
    assert custom.loc["x", "M"] == "1.000*"

    plain = hub.regression_table(stars=False)
    assert plain.loc["x", "M"] == "1.000"
    assert "Significance" not in plain.index


def test_native_docx_pdf_report_exports(tmp_path: Path) -> None:
    hub = OutputHub("Model Output Report", metadata={"project": "test"})
    hub.add_model(
        {
            "name": "M1",
            "params": {"x": 1.0},
            "std_errors": {"x": 0.1},
            "pvalues": {"x": 0.01},
            "statistics": {"N": 100},
        }
    )
    hub.add_table("Summary table", pd.DataFrame({"metric": ["N"], "value": [100]}))
    hub.add_note("Generated during automated testing.")

    docx_path = hub.export_report(tmp_path / "output_report.docx")
    pdf_path = hub.export_report(tmp_path / "output_report.pdf")

    assert docx_path.exists()
    assert pdf_path.exists()

    bundle = hub.export_bundle(tmp_path / "bundle")

    assert (tmp_path / "bundle" / "reports" / "output_report.docx").exists()
    assert (tmp_path / "bundle" / "reports" / "output_report.pdf").exists()
    assert "reports" in bundle


def test_comprehensive_output_exports(tmp_path: Path) -> None:
    hub = OutputHub("Model Output Report", metadata={"project": "test"})
    hub.add_model(
        {
            "name": "M1",
            "params": {"x": 1.0},
            "std_errors": {"x": 0.1},
            "pvalues": {"x": 0.01},
            "statistics": {"N": 100},
            "diagnostics": {"Hansen p": 0.50},
        }
    )
    hub.add_table("Summary table", pd.DataFrame({"metric": ["N"], "value": [100]}))
    hub.add_note("Generated during automated testing.")
    assert hub.export_report(tmp_path / "output_report.docx").exists()
    assert hub.export_report(tmp_path / "output_report.pdf").exists()
    assert hub.export_report(tmp_path / "output_report.tex").exists()
    assert hub.export_excel_workbook(tmp_path / "output_workbook.xlsx").exists()
    assert hub.export_regression_table(tmp_path / "regression_table.pdf").exists()
    assert hub.export_regression_table(tmp_path / "regression_table.tex").exists()
    assert hub.export_regression_table(tmp_path / "regression_table.xlsx").exists()
    table_paths = hub.export_tables(tmp_path / "tables", formats=("xlsx", "tex", "pdf"))
    assert len(table_paths) == 3
    assert all(path.exists() for path in table_paths)
    bundle = hub.export_bundle(tmp_path / "bundle")
    assert (tmp_path / "bundle" / "reports" / "output_report.docx").exists()
    assert (tmp_path / "bundle" / "reports" / "output_report.pdf").exists()
    assert (tmp_path / "bundle" / "reports" / "output_report.tex").exists()
    assert (tmp_path / "bundle" / "workbooks" / "output_workbook.xlsx").exists()
    assert (tmp_path / "bundle" / "regression_tables" / "regression_table.pdf").exists()
    assert "reports" in bundle
    assert "workbook" in bundle


def test_embedded_diagnostics_are_extracted_from_regression_table() -> None:
    hub = OutputHub("Embedded diagnostics")

    table = pd.DataFrame(
        {
            "term": [
                "L.y",
                "x",
                "N",
                "Instruments",
                "AR(1) p",
                "AR(2) p",
                "Hansen p",
                "Sargan p",
                "Diff-Hansen p",
                "Entity FE",
                "Time FE",
            ],
            "coef": [0.42, 0.11, 946, 42, 0.085, 0.316, 0.190, 0.098, 0.089, "Yes", "Yes"],
            "se": [0.10, 0.04, None, None, None, None, None, None, None, None, None],
            "pvalue": [0.001, 0.030, None, None, None, None, None, None, None, None, None],
        }
    )

    hub.add_model_table(table, name="System GMM")

    output = hub.regression_table()

    assert "L.y" in output.index
    assert "x" in output.index
    assert "N" in output.index
    assert "Instruments" in output.index
    assert "AR(1) p" in output.index
    assert "AR(2) p" in output.index
    assert "Hansen p" in output.index
    assert "Sargan p" in output.index
    assert "Diff-Hansen p" in output.index
    assert "Entity FE" in output.index
    assert "Time FE" in output.index

    assert output.loc["N", "System GMM"] == "946"
    assert output.loc["Instruments", "System GMM"] == "42"
    assert output.loc["Hansen p", "System GMM"] == "0.190"
    assert output.loc["Entity FE", "System GMM"] == "Yes"
    assert "***" in output.loc["L.y", "System GMM"]


def test_table_notes_are_added_below_regression_table(tmp_path: Path) -> None:
    hub = OutputHub("Table Notes Test")
    hub.add_model(
        {
            "name": "M1",
            "params": {"x": 1.0},
            "std_errors": {"x": 0.1},
            "pvalues": {"x": 0.01},
            "statistics": {"N": 100},
        }
    )
    hub.add_table_note("Standard errors in parentheses.")
    table = hub.regression_table()
    assert "Notes" in table.index


def test_table_notes_span_all_columns_in_rich_exports(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import load_workbook

    hub = OutputHub("Spanning Notes")
    for name in ("M1", "M2", "M3"):
        hub.add_model(
            {
                "name": name,
                "params": {"x": 1.0},
                "std_errors": {"x": 0.1},
                "pvalues": {"x": 0.01},
            }
        )
    note = "Standard errors in parentheses and clustered by entity."
    hub.add_table_note(note)

    html_path = hub.export_regression_table(tmp_path / "table.html")
    html = html_path.read_text(encoding="utf-8")
    assert 'colspan="4" class="table-note">Notes: ' + note in html

    latex_path = hub.export_regression_table(tmp_path / "table.tex")
    latex = latex_path.read_text(encoding="utf-8")
    assert rf"\multicolumn{{4}}{{l}}{{Notes: {note}}}" in latex

    excel_path = hub.export_regression_table(tmp_path / "table.xlsx")
    sheet = load_workbook(excel_path).active
    note_cell = next(cell for cell in sheet["A"] if cell.value == f"Notes: {note}")
    assert f"A{note_cell.row}:D{note_cell.row}" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}

    docx_path = hub.export_report(tmp_path / "report.docx")
    document = Document(docx_path)
    note_row = next(row for row in document.tables[0].rows if row.cells[0].text == f"Notes: {note}")
    assert len({cell._tc for cell in note_row.cells}) == 1
