"""Report, table, workbook, and document exporters for Universal Output Hub."""

from __future__ import annotations

import shutil
from collections.abc import Mapping, Sequence
from html import escape
from pathlib import Path
from typing import Any

import pandas as pd

from .formatters import safe_filename


def _display_frame(
    frame: pd.DataFrame,
    *,
    include_index: bool = True,
    index_name: str = "",
) -> pd.DataFrame:
    attrs = dict(frame.attrs)
    output = frame.copy()
    if include_index:
        output = output.reset_index()
        first = output.columns[0]
        if first == "index" or first is None:
            output = output.rename(columns={first: index_name})
    output = output.fillna("").astype(str)
    output.attrs.update(attrs)
    return output


def _spanning_rows(frame: pd.DataFrame) -> dict[int, str]:
    """Return zero-based data-row positions that should span all columns."""
    return {
        int(item["position"]): str(item["text"])
        for item in frame.attrs.get("spanning_rows", [])
        if "position" in item and "text" in item
    }


def _frame_to_html(frame: pd.DataFrame, *, include_index: bool = True, index_name: str = "") -> str:
    display = _display_frame(frame, include_index=include_index, index_name=index_name)
    spans = _spanning_rows(display)
    parts = ['<table border="0" class="dataframe">', "  <thead>", '    <tr style="text-align: right;">']
    parts.extend(f"      <th>{escape(str(column))}</th>" for column in display.columns)
    parts.extend(["    </tr>", "  </thead>", "  <tbody>"])
    for position, row in enumerate(display.itertuples(index=False, name=None)):
        parts.append("    <tr>")
        if position in spans:
            parts.append(
                f'      <td colspan="{len(display.columns)}" class="table-note">{escape(spans[position])}</td>'
            )
        else:
            parts.extend(f"      <td>{escape(str(value))}</td>" for value in row)
        parts.append("    </tr>")
    parts.extend(["  </tbody>", "</table>"])
    return "\n".join(parts)


def _frame_to_latex(frame: pd.DataFrame, *, include_index: bool = True, index_name: str = "") -> str:
    display = _display_frame(frame, include_index=include_index, index_name=index_name)
    spans = _spanning_rows(display)
    column_count = len(display.columns)
    lines = [rf"\begin{{tabular}}{{{'l' + 'c' * (column_count - 1)}}}", r"\toprule"]
    lines.append(" & ".join(_latex_escape(column) for column in display.columns) + r" \\")
    lines.append(r"\midrule")
    for position, row in enumerate(display.itertuples(index=False, name=None)):
        if position in spans:
            lines.append(rf"\multicolumn{{{column_count}}}{{l}}{{{_latex_escape(spans[position])}}} \\")
        else:
            lines.append(" & ".join(_latex_escape(value) for value in row) + r" \\")
    lines.extend([r"\bottomrule", r"\end{tabular}"])
    return "\n".join(lines)


def _merge_excel_spanning_rows(writer: Any, sheet_name: str, frame: pd.DataFrame, *, include_index: bool) -> None:
    spans = _spanning_rows(frame)
    if not spans:
        return
    sheet = writer.sheets[sheet_name]
    column_count = len(frame.columns) + (1 if include_index else 0)
    for position, note in spans.items():
        excel_row = position + 2  # one-based row plus header
        sheet.merge_cells(start_row=excel_row, start_column=1, end_row=excel_row, end_column=column_count)
        sheet.cell(excel_row, 1, note)
        sheet.cell(excel_row, 1).alignment = __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")


def _iter_frame_rows(frame: pd.DataFrame) -> list[list[str]]:
    display = frame.fillna("").astype(str)
    return [[str(column) for column in display.columns], *display.values.tolist()]


def _safe_sheet_name(name: str, used: set[str] | None = None) -> str:
    used = used if used is not None else set()
    cleaned = safe_filename(name).replace("_", " ")[:31].strip() or "Sheet"
    base = cleaned[:31]
    candidate = base
    counter = 1
    while candidate in used:
        suffix = f" {counter}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _latex_escape(value: Any) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def _make_pdf_table(frame: pd.DataFrame, styles: Any) -> Any:
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    rows = _iter_frame_rows(frame)
    paragraph_rows = [[Paragraph(escape(str(cell)), styles["BodyText"]) for cell in row] for row in rows]
    table = Table(paragraph_rows, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
    ]
    for position in _spanning_rows(frame):
        commands.append(("SPAN", (0, position + 1), (-1, position + 1)))
        commands.append(("ALIGN", (0, position + 1), (-1, position + 1), "LEFT"))
    table.setStyle(TableStyle(commands))
    return table


def export_frame_pdf(
    frame: pd.DataFrame,
    path: str | Path,
    *,
    title: str | None = None,
    include_index: bool = True,
    index_name: str = "",
) -> Path:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ImportError("PDF export requires reportlab. Install with: pip install reportlab") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    story: list[Any] = []
    if title:
        story.extend([Paragraph(escape(title), styles["Title"]), Spacer(1, 0.15 * inch)])
    story.append(_make_pdf_table(_display_frame(frame, include_index=include_index, index_name=index_name), styles))
    doc.build(story)
    return path


def export_frame(
    frame: pd.DataFrame,
    path: str | Path,
    *,
    fmt: str | None = None,
    include_index: bool = True,
    sheet_name: str = "table",
    title: str | None = None,
    index_name: str = "",
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fmt = (fmt or path.suffix.lstrip(".") or "csv").lower()
    if fmt == "csv":
        frame.to_csv(path, index=include_index)
    elif fmt in {"xlsx", "xls"}:
        safe_sheet = _safe_sheet_name(sheet_name)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            frame.to_excel(writer, index=include_index, sheet_name=safe_sheet)
            _merge_excel_spanning_rows(writer, safe_sheet, frame, include_index=include_index)
    elif fmt in {"md", "markdown"}:
        path.write_text(frame.to_markdown(index=include_index), encoding="utf-8")
    elif fmt in {"html", "htm"}:
        path.write_text(_frame_to_html(frame, include_index=include_index, index_name=index_name), encoding="utf-8")
    elif fmt in {"tex", "latex"}:
        path.write_text(_frame_to_latex(frame, include_index=include_index, index_name=index_name), encoding="utf-8")
    elif fmt == "json":
        output = frame.reset_index() if include_index else frame
        path.write_text(output.to_json(orient="records", indent=2), encoding="utf-8")
    elif fmt == "pdf":
        export_frame_pdf(frame, path, title=title, include_index=include_index, index_name=index_name)
    else:
        raise ValueError(f"Unsupported table export format: {fmt}")
    return path


def _add_docx_table(document: Any, frame: pd.DataFrame) -> None:
    rows = _iter_frame_rows(frame)
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for idx, value in enumerate(rows[0]):
        table.rows[0].cells[idx].text = value
    spans = _spanning_rows(frame)
    for position, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        if position in spans:
            merged = cells[0].merge(cells[-1])
            merged.text = spans[position]
            continue
        for idx, value in enumerate(row):
            cells[idx].text = value


def export_docx_report(
    hub: Any,
    path: str | Path,
    *,
    regression_kwargs: Mapping[str, Any] | None = None,
    include_metadata: bool = True,
) -> Path:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise ImportError("DOCX export requires python-docx. Install with: pip install python-docx") from exc
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(hub.title, level=0)
    if include_metadata and hub.metadata:
        document.add_heading("Metadata", level=1)
        for key, value in hub.metadata.items():
            document.add_paragraph(f"{key}: {value}")
    if hub.notes:
        document.add_heading("Notes", level=1)
        for note in hub.notes:
            document.add_paragraph(note, style="List Bullet")
    if hub.models:
        document.add_heading("Regression / Model Results", level=1)
        frame = _display_frame(hub.regression_table(**dict(regression_kwargs or {})), index_name="term")
        _add_docx_table(document, frame)
    if hub.tables:
        document.add_heading("Tables", level=1)
        for artifact in hub.tables:
            document.add_heading(artifact.name, level=2)
            if artifact.caption:
                document.add_paragraph(artifact.caption)
            _add_docx_table(document, _display_frame(artifact.data))
    if hub.figures:
        document.add_heading("Figures", level=1)
        for figure in hub.figures:
            document.add_heading(figure.name, level=2)
            if figure.caption:
                document.add_paragraph(figure.caption)
            try:
                document.add_picture(str(Path(figure.path)), width=Inches(6.0))
            except Exception as exc:
                document.add_paragraph(f"Figure could not be embedded: {figure.path} ({exc})")
    document.save(path)
    return path


def export_pdf_report(
    hub: Any,
    path: str | Path,
    *,
    regression_kwargs: Mapping[str, Any] | None = None,
    include_metadata: bool = True,
) -> Path:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ImportError("PDF export requires reportlab. Install with: pip install reportlab") from exc
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(escape(hub.title), styles["Title"]), Spacer(1, 0.2 * inch)]
    if include_metadata and hub.metadata:
        story.extend([Paragraph("Metadata", styles["Heading1"]), Spacer(1, 0.08 * inch)])
        for key, value in hub.metadata.items():
            story.append(Paragraph(f"<b>{escape(str(key))}</b>: {escape(str(value))}", styles["BodyText"]))
        story.append(Spacer(1, 0.15 * inch))
    if hub.notes:
        story.extend([Paragraph("Notes", styles["Heading1"]), Spacer(1, 0.08 * inch)])
        for note in hub.notes:
            story.append(Paragraph(f"- {escape(str(note))}", styles["BodyText"]))
        story.append(Spacer(1, 0.15 * inch))
    if hub.models:
        story.extend([Paragraph("Regression / Model Results", styles["Heading1"]), Spacer(1, 0.08 * inch)])
        frame = _display_frame(hub.regression_table(**dict(regression_kwargs or {})), index_name="term")
        story.append(_make_pdf_table(frame, styles))
        story.append(Spacer(1, 0.2 * inch))
    if hub.tables:
        story.extend([Paragraph("Tables", styles["Heading1"]), Spacer(1, 0.08 * inch)])
        for artifact in hub.tables:
            story.append(Paragraph(escape(artifact.name), styles["Heading2"]))
            if artifact.caption:
                story.append(Paragraph(escape(artifact.caption), styles["Italic"]))
            story.append(_make_pdf_table(_display_frame(artifact.data), styles))
            story.append(Spacer(1, 0.2 * inch))
    if hub.figures:
        story.extend([Paragraph("Figures", styles["Heading1"]), Spacer(1, 0.08 * inch)])
        for figure in hub.figures:
            story.append(Paragraph(escape(figure.name), styles["Heading2"]))
            if figure.caption:
                story.append(Paragraph(escape(figure.caption), styles["Italic"]))
            try:
                image = Image(str(Path(figure.path)))
                max_width = 6.7 * inch
                max_height = 8.5 * inch
                scale = min(max_width / image.drawWidth, max_height / image.drawHeight, 1.0)
                image.drawWidth *= scale
                image.drawHeight *= scale
                story.append(image)
            except Exception as exc:
                message = f"Figure could not be embedded: {figure.path} ({exc})"
                story.append(Paragraph(escape(message), styles["BodyText"]))
            story.append(Spacer(1, 0.2 * inch))
    doc.build(story)
    return path


def export_latex_report(
    hub: Any,
    path: str | Path,
    *,
    regression_kwargs: Mapping[str, Any] | None = None,
    include_metadata: bool = True,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    parts = [
        r"\documentclass[11pt,a4paper]{article}",
        r"\usepackage[margin=1in]{geometry}",
        r"\usepackage{booktabs}",
        r"\usepackage{graphicx}",
        r"\usepackage{longtable}",
        r"\usepackage{array}",
        r"\begin{document}",
        rf"\title{{{_latex_escape(hub.title)}}}",
        r"\date{}",
        r"\maketitle",
    ]
    if include_metadata and hub.metadata:
        parts.append(r"\section*{Metadata}")
        for key, value in hub.metadata.items():
            parts.append(rf"\textbf{{{_latex_escape(key)}}}: {_latex_escape(value)}\\")
    if hub.notes:
        parts.append(r"\section*{Notes}")
        parts.append(r"\begin{itemize}")
        for note in hub.notes:
            parts.append(rf"\item {_latex_escape(note)}")
        parts.append(r"\end{itemize}")
    if hub.models:
        parts.append(r"\section*{Regression / Model Results}")
        table = hub.regression_table(**dict(regression_kwargs or {}))
        parts.append(_frame_to_latex(table, include_index=True, index_name="term"))
    if hub.tables:
        parts.append(r"\section*{Tables}")
        for artifact in hub.tables:
            parts.append(rf"\subsection*{{{_latex_escape(artifact.name)}}}")
            if artifact.caption:
                parts.append(_latex_escape(artifact.caption))
            parts.append(artifact.data.to_latex(escape=False))
    if hub.figures:
        parts.append(r"\section*{Figures}")
        for figure in hub.figures:
            parts.append(rf"\subsection*{{{_latex_escape(figure.name)}}}")
            if figure.caption:
                parts.append(_latex_escape(figure.caption))
            figure_path = Path(figure.path).as_posix()
            parts.append(r"\begin{center}")
            parts.append(rf"\includegraphics[width=0.9\textwidth]{{{figure_path}}}")
            parts.append(r"\end{center}")
    parts.append(r"\end{document}")
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


def export_excel_workbook(
    hub: Any,
    path: str | Path,
    *,
    regression_kwargs: Mapping[str, Any] | None = None,
    include_metadata: bool = True,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    used_sheets: set[str] = set()
    wrote_sheet = False
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if hub.models:
            sheet = _safe_sheet_name("regression_table", used_sheets)
            regression = hub.regression_table(**dict(regression_kwargs or {}))
            regression.to_excel(writer, sheet_name=sheet)
            _merge_excel_spanning_rows(writer, sheet, regression, include_index=True)
            wrote_sheet = True
        for artifact in hub.tables:
            sheet = _safe_sheet_name(artifact.name, used_sheets)
            artifact.data.to_excel(writer, sheet_name=sheet, index=True)
            wrote_sheet = True
        if include_metadata and (hub.metadata or hub.notes or not wrote_sheet):
            rows = [{"field": key, "value": value} for key, value in hub.metadata.items()]
            rows.extend({"field": "note", "value": note} for note in hub.notes)
            if not rows:
                rows = [{"field": "status", "value": "No model or table outputs were registered."}]
            pd.DataFrame(rows).to_excel(writer, sheet_name=_safe_sheet_name("metadata", used_sheets), index=False)
    return path


def export_report(
    hub: Any,
    path: str | Path,
    *,
    fmt: str | None = None,
    regression_kwargs: Mapping[str, Any] | None = None,
    include_metadata: bool = True,
) -> Path:
    path = Path(path)
    fmt = (fmt or path.suffix.lstrip(".") or "docx").lower()
    if fmt == "docx":
        return export_docx_report(hub, path, regression_kwargs=regression_kwargs, include_metadata=include_metadata)
    if fmt == "pdf":
        return export_pdf_report(hub, path, regression_kwargs=regression_kwargs, include_metadata=include_metadata)
    if fmt in {"tex", "latex"}:
        return export_latex_report(hub, path, regression_kwargs=regression_kwargs, include_metadata=include_metadata)
    if fmt in {"xlsx", "xls"}:
        return export_excel_workbook(hub, path, regression_kwargs=regression_kwargs, include_metadata=include_metadata)
    if fmt in {"html", "htm"}:
        html = hub.export_html_report(path.parent, regression_kwargs=regression_kwargs)
        if html != path:
            shutil.copy2(html, path)
        return path
    raise ValueError(f"Unsupported report export format: {fmt}")


def attach_output_methods(output_hub_cls: type[Any]) -> None:
    if getattr(output_hub_cls, "_uoh_output_methods_attached", False):
        return

    def export_regression_table_method(self: Any, path: str | Path, *, fmt: str | None = None, **kwargs: Any) -> Path:
        table = self.regression_table(**kwargs)
        return export_frame(table, path, fmt=fmt, include_index=True, title="Regression table", index_name="term")

    def export_tables_method(
        self: Any,
        output_dir: str | Path,
        *,
        formats: Sequence[str] = ("csv", "xlsx", "html", "md", "tex", "pdf", "json"),
    ) -> list[Path]:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        written: list[Path] = []
        for artifact in self.tables:
            base = safe_filename(artifact.name)
            for table_format in formats:
                fmt = table_format.lower()
                suffix = "md" if fmt == "markdown" else ("tex" if fmt == "latex" else fmt)
                written.append(
                    export_frame(
                        artifact.data,
                        out_dir / f"{base}.{suffix}",
                        fmt=fmt,
                        include_index=True,
                        sheet_name=base[:31],
                        title=artifact.name,
                    )
                )
        return written

    def export_report_method(self: Any, path: str | Path, **kwargs: Any) -> Path:
        return export_report(self, path, **kwargs)

    def export_docx_report_method(self: Any, path: str | Path, **kwargs: Any) -> Path:
        return export_docx_report(self, path, **kwargs)

    def export_pdf_report_method(self: Any, path: str | Path, **kwargs: Any) -> Path:
        return export_pdf_report(self, path, **kwargs)

    def export_latex_report_method(self: Any, path: str | Path, **kwargs: Any) -> Path:
        return export_latex_report(self, path, **kwargs)

    def export_excel_workbook_method(self: Any, path: str | Path, **kwargs: Any) -> Path:
        return export_excel_workbook(self, path, **kwargs)

    def export_bundle_with_all_outputs(
        self: Any,
        output_dir: str | Path,
        *,
        regression_formats: Sequence[str] = ("csv", "xlsx", "html", "md", "tex", "pdf", "json"),
        table_formats: Sequence[str] = ("csv", "xlsx", "html", "md", "tex", "pdf", "json"),
        report_formats: Sequence[str] = ("docx", "pdf", "tex"),
        report_filename: str = "output_report",
        workbook_filename: str = "output_workbook",
        include_workbook: bool = True,
        regression_kwargs: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        from .core import FigureArtifact

        root = Path(output_dir)
        root.mkdir(parents=True, exist_ok=True)
        fig_dir = root / "figures"
        fig_dir.mkdir(exist_ok=True)
        copied_figures = []
        for fig in self.figures:
            src = Path(fig.path)
            dst = fig_dir / safe_filename(src.name)
            if src.exists() and src.resolve() != dst.resolve():
                shutil.copy2(src, dst)
            copied_figures.append(FigureArtifact(fig.name, str(dst), fig.caption, fig.metadata))
        if copied_figures:
            self.figures = copied_figures
        written_regression: list[Path] = []
        if self.models:
            reg_dir = root / "regression_tables"
            reg_dir.mkdir(exist_ok=True)
            for reg_format in regression_formats:
                fmt = reg_format.lower()
                suffix = "md" if fmt == "markdown" else ("tex" if fmt == "latex" else fmt)
                written_regression.append(
                    self.export_regression_table(
                        reg_dir / f"regression_table.{suffix}",
                        fmt=fmt,
                        **dict(regression_kwargs or {}),
                    )
                )
        table_paths = self.export_tables(root / "tables", formats=table_formats) if self.tables else []
        manifest = self.export_manifest(root)
        html = self.export_html_report(root, regression_kwargs=regression_kwargs)
        reports: list[Path] = []
        if report_formats:
            reports_dir = root / "reports"
            reports_dir.mkdir(exist_ok=True)
            for report_format in report_formats:
                fmt = report_format.lower()
                suffix = "tex" if fmt == "latex" else ("html" if fmt == "htm" else fmt)
                reports.append(
                    export_report(
                        self,
                        reports_dir / safe_filename(report_filename, suffix=suffix),
                        fmt=fmt,
                        regression_kwargs=regression_kwargs,
                    )
                )
        workbook = None
        if include_workbook:
            workbooks_dir = root / "workbooks"
            workbooks_dir.mkdir(exist_ok=True)
            workbook = export_excel_workbook(
                self,
                workbooks_dir / safe_filename(workbook_filename, suffix="xlsx"),
                regression_kwargs=regression_kwargs,
            )
        return {
            "root": root,
            "manifest": manifest,
            "html_report": html,
            "regression_tables": written_regression,
            "tables": table_paths,
            "figures": [Path(f.path) for f in self.figures],
            "reports": reports,
            "workbook": workbook,
        }

    output_hub_cls.export_regression_table = export_regression_table_method
    output_hub_cls.export_tables = export_tables_method
    output_hub_cls.export_report = export_report_method
    output_hub_cls.export_docx_report = export_docx_report_method
    output_hub_cls.export_pdf_report = export_pdf_report_method
    output_hub_cls.export_latex_report = export_latex_report_method
    output_hub_cls.export_excel_workbook = export_excel_workbook_method
    output_hub_cls.export_bundle = export_bundle_with_all_outputs
    output_hub_cls._uoh_output_methods_attached = True


attach_report_methods = attach_output_methods
