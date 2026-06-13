from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any

import pandas as pd

from universal_output_hub.normalizer import normalize_result
from universal_output_hub.schema import NormalizedCoefficient, NormalizedResult
from universal_output_hub.table import OutputTable

_DIAG_MAP = {
    "n": "nobs",
    "observations": "nobs",
    "obs": "nobs",
    "instruments": "n_instruments",
    "instrument count": "n_instruments",
    "hansen p": "hansen_p",
    "sargan p": "sargan_p",
    "diff-hansen p": "diff_hansen_p",
    "difference-in-hansen p": "diff_hansen_p",
    "ar(1)": "ar1",
    "ar(1) p": "ar1_p",
    "ar1 p": "ar1_p",
    "ar(2)": "ar2",
    "ar(2) p": "ar2_p",
    "ar2 p": "ar2_p",
}


def _is_missing(value: Any) -> bool:
    if value is None:
        return True

    try:
        result = pd.isna(value)
    except TypeError:
        return False

    if isinstance(result, bool):
        return result

    return False


def _clean_value(value: Any) -> Any:
    if _is_missing(value):
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _find_column(frame: pd.DataFrame, candidates: list[str]) -> str | None:
    columns = {str(col).lower(): str(col) for col in frame.columns}

    for candidate in candidates:
        if candidate.lower() in columns:
            return columns[candidate.lower()]

    return None


def _is_diagnostic_term(term: str) -> bool:
    key = term.strip().lower()
    return key in _DIAG_MAP or key in {
        "entity fe",
        "time fe",
        "fixed effects",
        "controls",
    }


class OutputHub:
    """Unified export facade for model outputs.

    Supports both APIs:

    - New style: OutputHub([result], model_names=["M1"])
    - Legacy style: OutputHub("Title").add_model(result)
    """

    def __init__(
        self,
        results_or_title: Any | None = None,
        *,
        results: list[Any] | None = None,
        model_names: list[str] | None = None,
        variable_labels: dict[str, str] | None = None,
        metadata: dict[str, Any] | None = None,
        show_stars: bool = True,
        star_levels: dict[float, str] | list[tuple[float, str]] | None = None,
        title: str | None = None,
    ) -> None:
        self.title: str | None = title
        self.metadata: dict[str, Any] = metadata or {}
        self.variable_labels = variable_labels or {}
        self.show_stars = show_stars
        self.star_levels = star_levels
        self.notes: list[str] = []
        self._results: list[NormalizedResult] = []

        initial_results: list[Any] = []

        if isinstance(results_or_title, str):
            self.title = results_or_title
        elif results_or_title is not None:
            if isinstance(results_or_title, list):
                initial_results.extend(results_or_title)
            else:
                initial_results.append(results_or_title)

        if results:
            initial_results.extend(results)

        if model_names is not None and len(model_names) != len(initial_results):
            raise ValueError("model_names must match the number of results.")

        for index, result in enumerate(initial_results):
            name = model_names[index] if model_names else None
            self.add_model(result, name=name)

    @property
    def results(self) -> list[NormalizedResult]:
        return list(self._results)

    def add_model(
        self,
        model: Any,
        *,
        name: str | None = None,
        estimator: str | None = None,
        backend: str | None = None,
    ) -> OutputHub:
        normalized = normalize_result(
            model,
            name=name,
            estimator=estimator,
            backend=backend,
        )

        diagnostics = dict(normalized.diagnostics)
        nobs = normalized.nobs
        n_instruments = normalized.n_instruments
        n_groups = normalized.n_groups

        if isinstance(model, Mapping):
            stats = model.get("statistics", {})
            if isinstance(stats, Mapping):
                for key, value in stats.items():
                    clean = _clean_value(value)
                    diagnostics[str(key)] = clean
                    mapped = _DIAG_MAP.get(str(key).strip().lower())
                    if mapped:
                        diagnostics[mapped] = clean

                if nobs is None and "N" in stats:
                    try:
                        nobs = int(stats["N"])
                    except (TypeError, ValueError):
                        pass

            extra = model.get("diagnostics", {})
            if isinstance(extra, Mapping):
                for key, value in extra.items():
                    clean = _clean_value(value)
                    diagnostics[str(key)] = clean
                    mapped = _DIAG_MAP.get(str(key).strip().lower())
                    if mapped:
                        diagnostics[mapped] = clean

        if nobs is not None:
            diagnostics.setdefault("nobs", nobs)
        if n_groups is not None:
            diagnostics.setdefault("n_groups", n_groups)
        if n_instruments is not None:
            diagnostics.setdefault("n_instruments", n_instruments)

        self._results.append(
            NormalizedResult(
                name=normalized.name,
                estimator=normalized.estimator,
                backend=normalized.backend,
                covariance_type=normalized.covariance_type,
                nobs=nobs,
                n_groups=n_groups,
                n_instruments=n_instruments,
                coefficients=normalized.coefficients,
                diagnostics=diagnostics,
                metadata=normalized.metadata,
            )
        )

        return self

    def add_model_table(self, table: pd.DataFrame, *, name: str = "model") -> OutputHub:
        term_col = _find_column(table, ["term", "variable", "name"])
        coef_col = _find_column(table, ["coef", "coefficient", "estimate", "params"])
        se_col = _find_column(table, ["se", "std_error", "std_errors", "stderr"])
        p_col = _find_column(table, ["pvalue", "p_value", "pvalues", "p"])

        if term_col is None or coef_col is None:
            raise ValueError("add_model_table requires term and coefficient columns.")

        coefficients: list[NormalizedCoefficient] = []
        diagnostics: dict[str, Any] = {}

        for _, row in table.iterrows():
            term = str(row[term_col])
            coef = _clean_value(row[coef_col])
            se = _clean_value(row[se_col]) if se_col else None
            p_value = _clean_value(row[p_col]) if p_col else None

            is_coefficient = not _is_diagnostic_term(term) and (
                se is not None or p_value is not None
            )

            if is_coefficient:
                coefficients.append(
                    NormalizedCoefficient(
                        name=term,
                        estimate=float(coef) if isinstance(coef, int | float) else None,
                        std_error=float(se) if isinstance(se, int | float) else None,
                        p_value=float(p_value) if isinstance(p_value, int | float) else None,
                    )
                )
                continue

            diagnostics[term] = coef
            mapped = _DIAG_MAP.get(term.strip().lower())
            if mapped:
                diagnostics[mapped] = coef

        self._results.append(
            NormalizedResult(
                name=name,
                coefficients=coefficients,
                diagnostics=diagnostics,
            )
        )

        return self

    def add_note(self, note: str) -> OutputHub:
        self.notes.append(str(note))
        return self

    def add_table_note(self, note: str) -> OutputHub:
        return self.add_note(note)

    def table(
        self,
        *,
        show_stars: bool | None = None,
        stars: bool | None = None,
        star_levels: dict[float, str] | list[tuple[float, str]] | None = None,
    ) -> OutputTable:
        if not self._results:
            raise ValueError("OutputHub has no models. Add a model before exporting.")

        if stars is not None:
            active_show_stars = stars
        elif show_stars is not None:
            active_show_stars = show_stars
        else:
            active_show_stars = self.show_stars

        return OutputTable(
            self._results,
            title=self.title,
            metadata=self.metadata,
            notes=self.notes,
            variable_labels=self.variable_labels,
            show_stars=active_show_stars,
            star_levels=star_levels or self.star_levels,
        )

    def to_markdown(
        self,
        path: str | Path | None = None,
        *,
        show_stars: bool | None = None,
        stars: bool | None = None,
        star_levels: dict[float, str] | list[tuple[float, str]] | None = None,
    ) -> str:
        content = self.table(
            show_stars=show_stars,
            stars=stars,
            star_levels=star_levels,
        ).to_markdown()

        if path is not None:
            output_path = Path(path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")

        return content

    def to_csv(self, path: str | Path | None = None) -> str | None:
        table = self.table()

        if path is None:
            return table.to_csv_string()

        table.to_csv(path)
        return None

    def to_latex(self, path: str | Path | None = None) -> str:
        content = self.table().to_latex()

        if path is not None:
            output_path = Path(path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")

        return content

    def to_html(self, path: str | Path | None = None) -> str:
        content = self.table().to_html()

        if path is not None:
            output_path = Path(path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")

        return content

    def to_diagnostics(self, path: str | Path | None = None) -> str:
        diagnostics = self.table().diagnostics_frame()

        if diagnostics.empty:
            content = "No diagnostics available.\n"
        else:
            content = diagnostics.to_markdown(index=False) + "\n"

        if path is not None:
            output_path = Path(path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")

        return content

    def to_docx(self, path: str | Path) -> None:
        output_path = Path(path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            from docx import Document
        except ImportError:
            output_path.write_text(self.to_markdown(), encoding="utf-8")
            return

        document = Document()

        if self.title:
            document.add_heading(self.title, level=1)

        if self.metadata:
            document.add_heading("Metadata", level=2)
            for key, value in self.metadata.items():
                document.add_paragraph(f"{key}: {value}")

        document.add_heading("Coefficients", level=2)
        coef = self.table().coefficient_frame()
        table = document.add_table(rows=1, cols=len(coef.columns))
        for i, column in enumerate(coef.columns):
            table.rows[0].cells[i].text = str(column)

        for _, row in coef.iterrows():
            cells = table.add_row().cells
            for i, column in enumerate(coef.columns):
                cells[i].text = str(row[column])

        diagnostics = self.table().diagnostics_frame()
        if not diagnostics.empty:
            document.add_heading("Diagnostics", level=2)
            for _, row in diagnostics.iterrows():
                document.add_paragraph(
                    " | ".join(str(row[col]) for col in diagnostics.columns)
                )

        if self.notes:
            document.add_heading("Notes", level=2)
            for note in self.notes:
                document.add_paragraph(note)

        document.save(output_path)

    def to_pdf(self, path: str | Path) -> None:
        output_path = Path(path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
        except ImportError:
            output_path.write_bytes(
                b"%PDF-1.4\n1 0 obj<<>>endobj\n"
                b"trailer<<>>\n%%EOF\n"
            )
            return

        pdf = canvas.Canvas(str(output_path), pagesize=letter)
        _, height = letter
        y = height - 40

        for line in self.to_markdown().splitlines():
            if y < 40:
                pdf.showPage()
                y = height - 40
            pdf.drawString(40, y, line[:110])
            y -= 14

        pdf.save()

    def to_bundle(self, directory: str | Path) -> None:
        output_dir = Path(directory)
        output_dir.mkdir(parents=True, exist_ok=True)

        self.to_markdown(output_dir / "results.md")
        self.to_csv(output_dir / "results.csv")
        self.to_latex(output_dir / "results.tex")
        self.to_html(output_dir / "results.html")
        self.to_diagnostics(output_dir / "diagnostics.md")
        self.to_docx(output_dir / "report.docx")
        self.to_pdf(output_dir / "report.pdf")

    render = to_markdown
    export_markdown = to_markdown
    save_markdown = to_markdown
    export_csv = to_csv
    save_csv = to_csv
    export_latex = to_latex
    save_latex = to_latex
    export_html = to_html
    save_html = to_html
    export_docx = to_docx
    save_docx = to_docx
    export_pdf = to_pdf
    save_pdf = to_pdf
    export_all = to_bundle
    save_all = to_bundle
