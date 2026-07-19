"""Core API for Universal Output Hub."""

from __future__ import annotations

import json
import re
import shutil
import warnings
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd

from .adapters import RegressionModel, from_coefficient_table, normalise_model
from .formatters import (
    as_float,
    format_number,
    format_significance_note,
    normalise_label,
    safe_filename,
    significance_stars,
)


@dataclass(slots=True)
class TableArtifact:
    name: str
    data: pd.DataFrame
    caption: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "caption": self.caption,
            "columns": [str(c) for c in self.data.columns],
            "rows": len(self.data),
            "metadata": self.metadata,
        }


@dataclass(slots=True)
class FigureArtifact:
    name: str
    path: str
    caption: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {"name": self.name, "path": self.path, "caption": self.caption, "metadata": self.metadata}


_TABLE_TEMPLATES: dict[str, dict[str, Any]] = {
    "economics": {
        "stats_order": [
            "N",
            "Groups",
            "Instruments",
            "Clusters",
            "Events",
            "Choice sets",
            "Alternatives",
            "Categories",
            "R2",
            "Adj. R2",
            "Within R2",
            "Entity FE",
            "Time FE",
            "Fixed effects",
            "Clustered SE",
            "Converged",
            "Inference valid",
            "AR(1) p",
            "AR(2) p",
            "Hansen p",
            "Sargan p",
            "Diff-Hansen p",
            "Backend",
            "Covariance type",
        ],
        "star_levels": {"***": 0.01, "**": 0.05, "*": 0.1},
        "star_note": "* p≤0.1, ** p≤0.05, *** p≤0.01",
    },
    "journal": {
        "stats_order": [
            "N",
            "R2",
            "Adj. R2",
            "Entity FE",
            "Time FE",
            "Clustered SE",
            "Instruments",
            "Hansen p",
            "AR(1) p",
            "AR(2) p",
        ],
        "star_levels": {"***": 0.01, "**": 0.05, "*": 0.1},
        "star_note": "* p≤0.1, ** p≤0.05, *** p≤0.01",
    },
    "stata": {
        "stats_order": [
            "N",
            "R2",
            "Adj. R2",
            "Entity FE",
            "Time FE",
            "Fixed effects",
            "Clustered SE",
            "Instruments",
            "Hansen p",
            "Sargan p",
            "AR(1) p",
            "AR(2) p",
        ],
        "star_levels": {"***": 0.01, "**": 0.05, "*": 0.1},
        "star_note": "* p≤0.1, ** p≤0.05, *** p≤0.01",
    },
}


def _resolve_table_template(template: str | None) -> dict[str, Any]:
    if template is None:
        return {}

    key = str(template).strip().lower()
    if key in {"", "default", "none"}:
        return {}

    if key not in _TABLE_TEMPLATES:
        valid = ", ".join(["default", *_TABLE_TEMPLATES])
        raise ValueError(f"Unknown table template '{template}'. Valid templates: {valid}.")

    return dict(_TABLE_TEMPLATES[key])


class OutputHub:
    """Collect, standardise, and export research outputs.

    The hub is deliberately model-agnostic. It accepts:
    - fitted model objects from statsmodels, linearmodels, pyfixest-like APIs;
    - plain dictionaries for custom estimators such as System GMM;
    - coefficient tables exported by Stata/R/SPSS/EViews/etc.;
    - pandas tables;
    - matplotlib-like figures or existing graph files.
    """

    def __init__(self, title: str = "Research Output Hub", *, metadata: Mapping[str, Any] | None = None) -> None:
        self.title = normalise_label(title)
        self.metadata: dict[str, Any] = dict(metadata or {})
        self.models: list[RegressionModel] = []
        self.tables: list[TableArtifact] = []
        self.figures: list[FigureArtifact] = []
        self.notes: list[str] = []
        self.table_notes: list[str] = []

    # ------------------------------------------------------------------
    # Add outputs
    # ------------------------------------------------------------------
    def add_model(
        self,
        result: Any,
        *,
        name: str | None = None,
        diagnostics: Mapping[str, Any] | None = None,
        adapter: str = "auto",
    ) -> RegressionModel:
        model = normalise_model(result, name=name, diagnostics=diagnostics, adapter=adapter)
        if model.params.empty:
            warnings.warn(
                f"Model '{model.name}' was added but no coefficients were extracted. "
                "Use add_model_table(...) or pass a mapping with params/std_errors/pvalues.",
                RuntimeWarning,
                stacklevel=2,
            )
        self.models.append(model)
        return model

    def add_model_table(
        self,
        table: pd.DataFrame,
        *,
        name: str,
        term_col: str = "term",
        coef_col: str = "coef",
        se_col: str | None = "se",
        pvalue_col: str | None = "pvalue",
        depvar: str | None = None,
        statistics: Mapping[str, Any] | None = None,
        diagnostics: Mapping[str, Any] | None = None,
        source: str = "external-table",
    ) -> RegressionModel:
        model = from_coefficient_table(
            table,
            name=name,
            term_col=term_col,
            coef_col=coef_col,
            se_col=se_col,
            pvalue_col=pvalue_col,
            depvar=depvar,
            statistics=statistics,
            diagnostics=diagnostics,
            source=source,
        )
        self.models.append(model)
        return model

    def add_model_file(
        self,
        path: str | Path,
        *,
        name: str,
        term_col: str = "term",
        coef_col: str = "coef",
        se_col: str | None = "se",
        pvalue_col: str | None = "pvalue",
        sheet_name: str | int | None = 0,
        depvar: str | None = None,
        statistics: Mapping[str, Any] | None = None,
        diagnostics: Mapping[str, Any] | None = None,
        source: str | None = None,
    ) -> RegressionModel:
        """Add a coefficient table from CSV, TSV, Excel, Stata .dta, or Parquet."""
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(p)
        suffix = p.suffix.lower()
        if suffix == ".csv":
            table = pd.read_csv(p)
        elif suffix in {".tsv", ".txt"}:
            table = pd.read_csv(p, sep="\t")
        elif suffix in {".xlsx", ".xls"}:
            table = pd.read_excel(p, sheet_name=sheet_name)
        elif suffix == ".dta":
            table = pd.read_stata(p)
        elif suffix in {".parquet", ".pq"}:
            table = pd.read_parquet(p)
        else:
            raise ValueError(f"Unsupported model file format: {suffix}")
        return self.add_model_table(
            table,
            name=name,
            term_col=term_col,
            coef_col=coef_col,
            se_col=se_col,
            pvalue_col=pvalue_col,
            depvar=depvar,
            statistics=statistics,
            diagnostics=diagnostics,
            source=source or f"file:{suffix}",
        )

    def add_table(
        self,
        name: str,
        table: pd.DataFrame | pd.Series | Mapping[str, Any] | Sequence[Any],
        *,
        caption: str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> TableArtifact:
        clean_name = normalise_label(name)
        if isinstance(table, pd.DataFrame):
            df = table.copy()
        elif isinstance(table, pd.Series):
            df = table.to_frame()
        elif isinstance(table, Mapping):
            df = pd.DataFrame.from_dict(table, orient="index", columns=["value"])
        else:
            df = pd.DataFrame(table)
        artifact = TableArtifact(clean_name, df, caption=caption, metadata=dict(metadata or {}))
        self.tables.append(artifact)
        return artifact

    def add_table_file(
        self,
        path: str | Path,
        *,
        name: str | None = None,
        caption: str | None = None,
        sheet_name: str | int | None = 0,
        metadata: Mapping[str, Any] | None = None,
    ) -> TableArtifact:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(p)
        suffix = p.suffix.lower()
        if suffix == ".csv":
            df = pd.read_csv(p)
        elif suffix in {".tsv", ".txt"}:
            df = pd.read_csv(p, sep="\t")
        elif suffix in {".xlsx", ".xls"}:
            df = pd.read_excel(p, sheet_name=sheet_name)
        elif suffix == ".dta":
            df = pd.read_stata(p)
        elif suffix in {".parquet", ".pq"}:
            df = pd.read_parquet(p)
        else:
            raise ValueError(f"Unsupported table file format: {suffix}")
        return self.add_table(name or p.stem, df, caption=caption, metadata=metadata)

    def add_figure(
        self,
        name: str,
        *,
        fig: Any | None = None,
        path: str | Path | None = None,
        output_dir: str | Path | None = None,
        caption: str | None = None,
        filename: str | None = None,
        dpi: int = 300,
        metadata: Mapping[str, Any] | None = None,
    ) -> FigureArtifact:
        clean_name = normalise_label(name)
        if fig is None and path is None:
            raise ValueError("Provide either fig=... or path=...")

        if fig is not None:
            if output_dir is None:
                raise ValueError("output_dir is required when saving a figure object.")
            out_dir = Path(output_dir)
            out_dir.mkdir(parents=True, exist_ok=True)
            fig_path = out_dir / safe_filename(filename or clean_name, suffix=".png")
            if not hasattr(fig, "savefig"):
                raise TypeError("fig must be a matplotlib-like object with savefig().")
            fig.savefig(fig_path, dpi=dpi, bbox_inches="tight")
            saved_path = fig_path
        else:
            source = Path(path)  # type: ignore[arg-type]
            if not source.exists():
                raise FileNotFoundError(source)
            if output_dir is None:
                saved_path = source
            else:
                out_dir = Path(output_dir)
                out_dir.mkdir(parents=True, exist_ok=True)
                saved_path = out_dir / safe_filename(filename or source.name)
                shutil.copy2(source, saved_path)

        artifact = FigureArtifact(clean_name, str(saved_path), caption=caption, metadata=dict(metadata or {}))
        self.figures.append(artifact)
        return artifact

    def add_note(self, note: str) -> None:
        self.notes.append(str(note).strip())

    def add_table_note(self, note: str) -> None:
        """Add an outreg-style note displayed below regression/model tables."""
        cleaned = str(note).strip()
        if cleaned:
            self.table_notes.append(cleaned)

    # ------------------------------------------------------------------
    # Regression-table construction
    # ------------------------------------------------------------------
    def regression_table(
        self,
        *,
        order: Sequence[str] | None = None,
        keep: Sequence[str] | None = None,
        drop: Sequence[str] | None = None,
        labels: Mapping[str, str] | None = None,
        stats_order: Sequence[str] | None = None,
        decimals: int = 3,
        stars: bool = True,
        star_levels: Mapping[str, float] | None = None,
        template: str | None = None,
        se_below: bool = True,
        include_depvar: bool = False,
        add_star_note: bool = True,
        star_note: str | None = None,
        table_notes: Sequence[str] | None = None,
    ) -> pd.DataFrame:
        if not self.models:
            raise ValueError("No models have been added.")

        labels = dict(labels or {})

        template_settings = _resolve_table_template(template)
        if template_settings:
            if stats_order is None:
                stats_order = template_settings.get("stats_order")
            if star_levels is None:
                star_levels = template_settings.get("star_levels")
            if star_note is None:
                star_note = template_settings.get("star_note")

        terms: list[str] = []
        for model in self.models:
            for term in model.terms:
                if term not in terms:
                    terms.append(term)

        if order:
            ordered = [term for term in order if term in terms]
            ordered.extend([term for term in terms if term not in ordered])
        else:
            ordered = terms

        if keep:
            patterns = [re.compile(p) for p in keep]
            ordered = [term for term in ordered if any(p.search(term) for p in patterns)]
        if drop:
            patterns = [re.compile(p) for p in drop]
            ordered = [term for term in ordered if not any(p.search(term) for p in patterns)]

        rows: list[tuple[str, list[str]]] = []
        if include_depvar:
            rows.append(("Dependent variable", [m.depvar or "" for m in self.models]))

        for term in ordered:
            coef_row: list[str] = []
            se_row: list[str] = []
            for model in self.models:
                coef = format_number(model.coefficient(term), decimals)
                if coef and stars:
                    coef += significance_stars(model.pvalue(term), star_levels)
                coef_row.append(coef)
                se = format_number(model.standard_error(term), decimals)
                se_row.append(f"({se})" if se else "")
            rows.append((labels.get(term, term), coef_row))
            if se_below:
                rows.append(("", se_row))

        default_stats = [
            "N",
            "Observed N",
            "Groups",
            "Instruments",
            "Clusters",
            "Events",
            "Choice sets",
            "Alternatives",
            "Categories",
            "R2",
            "Adj. R2",
            "Within R2",
            "Overall R2",
            "AIC",
            "BIC",
            "Log Likelihood",
            "F statistic",
            "F p-value",
            "AR(1)",
            "AR(1) p",
            "AR(2)",
            "AR(2) p",
            "Hansen p",
            "Sargan p",
            "Diff-Hansen p",
            "Backend",
            "Covariance type",
            "Entity FE",
            "Time FE",
            "Fixed effects",
            "Clustered SE",
            "Converged",
            "Inference valid",
        ]
        stat_keys = list(stats_order or default_stats)
        present_stats = [key for key in stat_keys if any(m.stat(key) is not None for m in self.models)]
        if present_stats:
            rows.append(("", ["" for _ in self.models]))
            for key in present_stats:
                stat_row = []
                for model in self.models:
                    value = model.stat(key)
                    number = as_float(value)
                    if isinstance(value, bool):
                        stat_row.append("Yes" if value else "No")
                    elif (
                        key
                        in {
                            "N",
                            "Observed N",
                            "Groups",
                            "Instruments",
                            "Clusters",
                            "Events",
                            "Choice sets",
                            "Alternatives",
                            "Categories",
                        }
                        and number is not None
                    ):
                        stat_row.append(str(int(round(number))))
                    elif number is not None:
                        stat_row.append(format_number(number, decimals))
                    elif value is not None:
                        stat_row.append(str(value))
                    else:
                        stat_row.append("")
                rows.append((key, stat_row))

        if stars and add_star_note:
            note = star_note or format_significance_note(star_levels)
            if note:
                rows.append(("", ["" for _ in self.models]))
                rows.append(("Significance", [note] + ["" for _ in self.models[1:]]))

        final_table_notes = list(self.table_notes)
        if table_notes:
            final_table_notes.extend(str(note).strip() for note in table_notes if str(note).strip())

        if final_table_notes:
            rows.append(("", ["" for _ in self.models]))
            for idx, note in enumerate(final_table_notes):
                label = "Notes" if idx == 0 else ""
                rows.append((label, [note] + ["" for _ in self.models[1:]]))

        df = pd.DataFrame({"term": [r[0] for r in rows]})
        for idx, model in enumerate(self.models):
            df[model.name] = [r[1][idx] for r in rows]
        result = df.set_index("term")

        # A DataFrame cannot express merged cells itself. Preserve the row
        # positions and complete note text so rich exporters can span the
        # entire displayed table instead of widening one model column.
        spanning_rows: list[dict[str, Any]] = []
        for position, (label, values) in enumerate(rows):
            if label in {"Significance", "Notes"} or (
                not label and values and values[0] and spanning_rows and spanning_rows[-1]["kind"] == "note"
            ):
                kind = "significance" if label == "Significance" else "note"
                prefix = "Significance: " if label == "Significance" else ("Notes: " if label == "Notes" else "")
                spanning_rows.append({"position": position, "text": prefix + values[0], "kind": kind})
        result.attrs["spanning_rows"] = spanning_rows
        return result

    # ------------------------------------------------------------------
    # Exporters
    # ------------------------------------------------------------------
    def export_regression_table(self, path: str | Path, *, fmt: str | None = None, **kwargs: Any) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        fmt = (fmt or path.suffix.lstrip(".") or "csv").lower()
        table = self.regression_table(**kwargs)
        if fmt == "csv":
            table.to_csv(path)
        elif fmt in {"xlsx", "xls"}:
            table.to_excel(path, sheet_name="regression_table")
        elif fmt in {"md", "markdown"}:
            path.write_text(table.to_markdown(), encoding="utf-8")
        elif fmt in {"html", "htm"}:
            path.write_text(table.to_html(border=0), encoding="utf-8")
        elif fmt in {"tex", "latex"}:
            path.write_text(table.to_latex(escape=False, column_format="l" + "c" * len(self.models)), encoding="utf-8")
        elif fmt == "json":
            path.write_text(table.reset_index().to_json(orient="records", indent=2), encoding="utf-8")
        else:
            raise ValueError(f"Unsupported regression table export format: {fmt}")
        return path

    def export_tables(
        self,
        output_dir: str | Path,
        *,
        formats: Sequence[str] = ("csv", "xlsx", "html", "md"),
    ) -> list[Path]:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        written: list[Path] = []
        for artifact in self.tables:
            base = safe_filename(artifact.name)
            for fmt in formats:
                fmt = fmt.lower()
                p = out_dir / f"{base}.{fmt}"
                if fmt == "csv":
                    artifact.data.to_csv(p, index=True)
                elif fmt in {"xlsx", "xls"}:
                    artifact.data.to_excel(p, index=True, sheet_name=base[:31])
                elif fmt in {"html", "htm"}:
                    p.write_text(artifact.data.to_html(border=0), encoding="utf-8")
                elif fmt in {"md", "markdown"}:
                    p.write_text(artifact.data.to_markdown(index=True), encoding="utf-8")
                elif fmt == "json":
                    p.write_text(artifact.data.to_json(orient="records", indent=2), encoding="utf-8")
                else:
                    raise ValueError(f"Unsupported table export format: {fmt}")
                written.append(p)
        return written

    def export_manifest(self, output_dir: str | Path) -> Path:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        manifest = {
            "title": self.title,
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "metadata": self.metadata,
            "models": [m.to_dict() for m in self.models],
            "tables": [t.to_dict() for t in self.tables],
            "figures": [f.to_dict() for f in self.figures],
            "notes": self.notes,
            "table_notes": self.table_notes,
        }
        path = out_dir / "manifest.json"
        path.write_text(json.dumps(manifest, indent=2, default=str), encoding="utf-8")
        return path

    def export_html_report(self, output_dir: str | Path, *, regression_kwargs: Mapping[str, Any] | None = None) -> Path:
        from .reports import _frame_to_html

        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        parts = [
            "<!doctype html>",
            "<html><head><meta charset='utf-8'>",
            f"<title>{self.title}</title>",
            "<style>"
            "body{font-family:Arial,sans-serif;margin:32px;line-height:1.45}"
            "table{border-collapse:collapse;margin:16px 0;width:auto}"
            "th,td{border:1px solid #ddd;padding:6px 10px;text-align:right}"
            "th:first-child,td:first-child{text-align:left}"
            "h1,h2{margin-top:28px}"
            "img{max-width:100%;height:auto;border:1px solid #eee}"
            "</style>",
            "</head><body>",
            f"<h1>{self.title}</h1>",
        ]
        if self.notes:
            parts.append("<h2>Notes</h2><ul>")
            for note in self.notes:
                parts.append(f"<li>{note}</li>")
            parts.append("</ul>")
        if self.models:
            parts.append("<h2>Regression / Model Results</h2>")
            table = self.regression_table(**dict(regression_kwargs or {}))
            parts.append(_frame_to_html(table, include_index=True, index_name="term"))
        if self.tables:
            parts.append("<h2>Tables</h2>")
            for artifact in self.tables:
                parts.append(f"<h3>{artifact.name}</h3>")
                if artifact.caption:
                    parts.append(f"<p><em>{artifact.caption}</em></p>")
                parts.append(artifact.data.to_html(border=0))
        if self.figures:
            parts.append("<h2>Figures</h2>")
            for fig in self.figures:
                rel = Path(fig.path).name
                parts.append(f"<h3>{fig.name}</h3>")
                if fig.caption:
                    parts.append(f"<p><em>{fig.caption}</em></p>")
                parts.append(f"<img src='figures/{rel}' alt='{fig.name}'>")
        parts.append("</body></html>")
        path = out_dir / "index.html"
        path.write_text("\n".join(parts), encoding="utf-8")
        return path

    def export_bundle(
        self,
        output_dir: str | Path,
        *,
        regression_formats: Sequence[str] = ("csv", "xlsx", "html", "md", "tex", "json"),
        table_formats: Sequence[str] = ("csv", "xlsx", "html", "md", "json"),
        regression_kwargs: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Export all registered outputs into a reproducible folder."""
        root = Path(output_dir)
        root.mkdir(parents=True, exist_ok=True)
        fig_dir = root / "figures"
        fig_dir.mkdir(exist_ok=True)

        # Copy externally registered figures into the bundle's figure folder.
        copied_figures: list[FigureArtifact] = []
        for fig in self.figures:
            src = Path(fig.path)
            dst = fig_dir / safe_filename(src.name)
            if src.exists() and src.resolve() != dst.resolve():
                shutil.copy2(src, dst)
            copied_figures.append(FigureArtifact(fig.name, str(dst), fig.caption, fig.metadata))
        self.figures = copied_figures

        written_regression: list[Path] = []
        if self.models:
            reg_dir = root / "regression_tables"
            reg_dir.mkdir(exist_ok=True)
            for fmt in regression_formats:
                suffix = "md" if fmt == "markdown" else fmt
                written_regression.append(
                    self.export_regression_table(
                        reg_dir / f"regression_table.{suffix}",
                        fmt=fmt,
                        **dict(regression_kwargs or {}),
                    )
                )

        table_paths = self.export_tables(root / "tables", formats=table_formats) if self.tables else []
        manifest = self.export_manifest(root)
        html = self.export_html_report(root, regression_kwargs=regression_kwargs)
        return {
            "root": root,
            "manifest": manifest,
            "html_report": html,
            "regression_tables": written_regression,
            "tables": table_paths,
            "figures": [Path(f.path) for f in self.figures],
        }


def _normalise_models_input(models: Any) -> list[Any]:
    """Return a list of model/result objects from a single object or sequence."""
    if isinstance(models, (str, bytes, Mapping)):
        return [models]
    if isinstance(models, Sequence) and not isinstance(models, (pd.DataFrame, pd.Series)):
        return list(models)
    return [models]


def _write_outreg_table(table: pd.DataFrame, path: Path) -> None:
    """Write an outreg-style table based on file extension."""
    from .reports import (
        _add_docx_table,
        _display_frame,
        _frame_to_html,
        _frame_to_latex,
        _make_pdf_table,
        _merge_excel_spanning_rows,
    )

    suffix = path.suffix.lower()

    if suffix == ".csv":
        table.to_csv(path)
        return

    if suffix in {".xlsx", ".xls"}:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            table.to_excel(writer, sheet_name="regression_table")
            _merge_excel_spanning_rows(writer, "regression_table", table, include_index=True)
        return

    if suffix in {".md", ".markdown"}:
        try:
            content = table.to_markdown()
        except Exception:
            content = table.to_string()
        path.write_text(content + "\n", encoding="utf-8")
        return

    if suffix in {".html", ".htm"}:
        path.write_text(_frame_to_html(table, include_index=True, index_name="term"), encoding="utf-8")
        return

    if suffix in {".tex", ".latex"}:
        path.write_text(_frame_to_latex(table, include_index=True, index_name="term"), encoding="utf-8")
        return

    if suffix == ".json":
        path.write_text(table.to_json(orient="split", indent=2), encoding="utf-8")
        return

    if suffix == ".txt":
        path.write_text(table.to_string() + "\n", encoding="utf-8")
        return

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError("DOCX export requires python-docx.") from exc

        doc = Document()
        doc.add_heading("Regression Results", level=1)

        _add_docx_table(doc, _display_frame(table, include_index=True, index_name="term"))

        doc.save(path)
        return

    if suffix == ".pdf":
        try:
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate
        except ImportError as exc:
            raise ImportError("PDF export requires reportlab.") from exc

        doc = SimpleDocTemplate(str(path), pagesize=landscape(letter))
        pdf_table = _make_pdf_table(_display_frame(table, include_index=True, index_name="term"), getSampleStyleSheet())
        doc.build([pdf_table])
        return

    raise ValueError(
        f"Unsupported outreg output format '{suffix}'. "
        "Use one of: .csv, .xlsx, .md, .html, .tex, .json, .txt, .docx, .pdf."
    )


def outreg(
    models: Any,
    using: str | Path,
    *,
    model_names: Sequence[str] | None = None,
    title: str = "Regression Results",
    template: str | None = None,
    stats: Sequence[str] | None = None,
    labels: Mapping[str, str] | None = None,
    notes: Sequence[str] | None = None,
    decimals: int = 3,
    stars: bool = True,
    star_levels: Mapping[str, float] | None = None,
    replace: bool = False,
    append: bool = False,
    adapter: str = "auto",
) -> Path:
    """Export outreg2-style side-by-side model results in one call.

    Parameters
    ----------
    models:
        A single fitted result/model object or a sequence of result/model objects.
    using:
        Output path. The file extension determines the export format.
    model_names:
        Optional display names for the models.
    title:
        Report title used internally by OutputHub.
    stats:
        Optional statistics/diagnostics row order.
    labels:
        Optional coefficient label mapping.
    notes:
        Optional notes appended below the table.
    decimals:
        Number of decimal places.
    stars:
        Whether to add significance stars.
    star_levels:
        Optional significance-star thresholds.
    replace:
        If False, refuse to overwrite an existing file.
    append:
        Reserved for future Stata-like append workflows. Currently not supported.
    adapter:
        Adapter mode passed to OutputHub.add_model().
    """
    if append:
        raise NotImplementedError("append=True is planned but not implemented in this release.")

    output_path = Path(using)
    if output_path.exists() and not replace:
        raise FileExistsError(f"{output_path} already exists. Pass replace=True to overwrite it.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    model_list = _normalise_models_input(models)
    names = list(model_names or [])

    if names and len(names) != len(model_list):
        raise ValueError("model_names must have the same length as models.")

    hub = OutputHub(title)

    for idx, model in enumerate(model_list):
        model_name = names[idx] if names else None
        hub.add_model(model, name=model_name, adapter=adapter)

    table = hub.regression_table(
        labels=labels,
        stats_order=stats,
        template=template,
        decimals=decimals,
        stars=stars,
        star_levels=star_levels,
        table_notes=notes,
    )

    _write_outreg_table(table, output_path)
    return output_path
